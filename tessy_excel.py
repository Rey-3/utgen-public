from openpyxl import load_workbook

### need to install https://openpyxl.readthedocs.io/en/stable/
def updateTestlistToExcel(func , templateFile , saveFileName , testList , testMethodlist , stubList , arg_name , arg_direction , arg_format):
    wb = load_workbook(templateFile)
    wsTC = wb.worksheets[1] ##Properties
    ws = wb.worksheets[5]   ##Values 
    wsStub = wb.worksheets[3] ##Stubs
    ### write Value sheet header (input/argument variable's name)
    typed_vars = set()      ### show the type only on the first column of each variable (RTE struct shown once)
    for j in range(len(arg_name)):
        base = arg_name[j].split('.')[0]
        first_of_var = base not in typed_vars
        typed_vars.add(base)
        ws.cell(row=1,column=2+j).value = arg_name[j]
        ws.cell(row=2,column=2+j).value = convertDirectionToTessyFormat(arg_direction[j] , arg_name[j] , arg_format[j] , first_of_var)
  
    ### write Value sheet content (test cases)
    ### write Properties sheet content (test design)
    for i in range(len(testList)):
        finalTestMethod = "Ignore"
        wsTC.cell(row=i+5,column=1).value = str(i+1)
        wsTC.cell(row=i+5,column=2).value = '0'
        
        for k in range(len(testMethodlist[i])):
            if testMethodlist[i][k] != finalTestMethod:
                if testMethodlist[i][k] == "Invalid Test Case":
                    finalTestMethod = "Invalid Test Case"
                    break
                elif testMethodlist[i][k] == "Error Injection Test":
                    finalTestMethod = "Error Injection Test"
                    break
                elif testMethodlist[i][k] == "Boundary Test" and finalTestMethod!="Error Injection Test":
                    finalTestMethod = "Boundary Test"
                elif testMethodlist[i][k] == "Functional Test" and finalTestMethod!= "Boundary Test" and finalTestMethod!="Error Injection Test":
                    finalTestMethod = "Functional Test"
        wsTC.cell(row=i+5,column=5).value = constructTestDescription(i+1, finalTestMethod , func )
        ws.cell(row=i+3,column=1).value = 'tc' + str(i+1) + '.1'
        for j in range(len(testList[i])):
            ws.cell(row=i+3,column=2+j).value = str(testList[i][j])
    ### write Stub sheet
    ### only stubs whose return value is actually supplied by this test (its var appears in the Value sheet)
    ### return that injected value; all others just return a zeroed temp.
    value_vars = {n.split('.')[0] for n in arg_name}
    for i in range(len(stubList)):
        name, ret_type, ret_name, _ = stubList[i]
        has_return = ret_type not in ("void", "static void")
        provides_value = ret_name in value_vars
        wsStub.cell(row=i+5,column=1).value = '0'
        wsStub.cell(row=i+5,column=2).value = '0'
        wsStub.cell(row=i+5,column=3).value = name
        stubBody = 'extern ' + ret_type + ' ' + ret_name + ';\r' if (has_return and provides_value) else ''
        if has_return:
            stubBody += ('return ' + ret_name + ';\r') if provides_value else (ret_type + ' temp = {0};\rreturn temp;\r')
        wsStub.cell(row=i+5,column=4).value = stubBody
    newFileName = '.\\generated\\' + func + '.xlsx'
    wb.save( saveFileName ) ##newFileName )

    
def constructTestDescription(tcNum , tcMtd , funcName):
    SourceName = '<TBD>'
    result = "Test Case ID: UT_" + SourceName + "_" + str(tcNum) + "\n"
    result += "Test Case Type: " + tcMtd + "\n"
    result += "Overview: 测试" + funcName  + "() 的正向测试"
    return result
   
def convertDirectionToTessyFormat(direction , name='' , fmt='', show_type=True):
    newDir = 'i' if direction == "in" else 'o'
    if name.startswith('tessy_') and show_type:   ### tessy_<name> variable shows its type once: i(SomeType); other members just i/o
        newDir += '(' + fmt + ')'
    return newDir
