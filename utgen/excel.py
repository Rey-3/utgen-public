from __future__ import annotations

import os
import re
import tempfile
import warnings
from collections.abc import Sequence
from copy import copy
from dataclasses import dataclass
from itertools import product
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DEFAULT_TEMPLATE_PATH = Path(r"D:\TestTeamAssets\common_tools\v2\template.xlsx")
XLSX_VARIABLE_LIMIT = 14_000
XLS_VARIABLES_PER_SHEET = 98
XLSX_ROW_LIMIT = 1_048_576
_REQUIRED_SHEETS = ("General", "Properties", "User Code", "Stubs", "Call Trace", "Values", "Fault Injections")
_DIRECTIONS = {"in", "out", "inout"}
_C_IDENTIFIER = re.compile(r"[A-Za-z_][A-Za-z0-9_]*")
_DIMENSION = re.compile(r"\[([^]]*)\]")
_C_BODY_FORBIDDEN = re.compile(r"[\r\n;{}]")
_STORAGE_QUALIFIERS = {"auto", "extern", "register", "static", "typedef", "_Thread_local"}


@dataclass(frozen=True)
class _ValueArgument:
    name: str
    length: str
    fmt: str
    direction: str
    pointer: bool
    values: tuple[str, ...]


@dataclass(frozen=True)
class _ValueSpec:
    path: Path
    function: str
    arguments: tuple[_ValueArgument, ...]
    row_count: int


@dataclass(frozen=True)
class _StubSpec:
    name: str
    return_type: str
    return_name: str
    argument_types: tuple[str, ...]


@dataclass(frozen=True)
class _Column:
    name: str
    declaration: str
    values: tuple[str, ...]


def _attribute(element, name, path):
    value = element.get(name)
    if value is None or not value.strip():
        raise ValueError(f"{path}: missing attribute {name}")
    return value.strip()


def _parse_value(path: Path) -> _ValueSpec:
    try:
        root = ET.parse(path).getroot()
    except (ET.ParseError, OSError) as error:
        raise ValueError(f"{path}: cannot parse value XML: {error}") from error
    if root.tag != "function":
        raise ValueError(f"{path}: root tag must be function")
    function = (root.get("name") or "").strip()
    if not function:
        raise ValueError(f"{path}: function name is required")
    if not _C_IDENTIFIER.fullmatch(function):
        raise ValueError(f"{path}: function name must be a valid C identifier")
    arguments = []
    names = set()
    row_count = None
    for element in root:
        if element.tag != "argument":
            raise ValueError(f"{path}: direct child must be argument")
        name = _attribute(element, "name", path)
        if name in names:
            raise ValueError(f"{path}: duplicate argument {name}")
        names.add(name)
        length = _attribute(element, "len", path)
        fmt = _attribute(element, "format", path)
        direction = _attribute(element, "direction", path)
        pointer_text = _attribute(element, "pointer", path)
        if direction not in _DIRECTIONS:
            raise ValueError(f"{path}: invalid direction {direction}")
        if pointer_text not in ("true", "false"):
            raise ValueError(f"{path}: invalid pointer {pointer_text}")
        values = ()
        if direction in ("in", "inout"):
            if element.get("testmethod") != "decision":
                raise ValueError(f"{path}: {name} requires testmethod=decision")
            text = element.get("range")
            if text is None or not text.strip():
                raise ValueError(f"{path}: {name} requires a non-empty range")
            parsed = []
            for token in text.split(","):
                token = token.strip()
                if not token:
                    raise ValueError(f"{path}: {name} contains an empty range token")
                parsed.append(token)
            values = tuple(parsed)
            if row_count is None:
                row_count = len(values)
            elif len(values) != row_count:
                raise ValueError(f"{path}: {name} range has {len(values)} rows; expected {row_count}")
        arguments.append(_ValueArgument(name, length, fmt, direction,
                                        pointer_text == "true", values))
    return _ValueSpec(path, function, tuple(arguments), row_count or 1)


def _stub_text(element, path: Path, label: str) -> str:
    text = element.text
    if text is None or not text.strip():
        raise ValueError(f"{path}: {label} is required")
    if _C_BODY_FORBIDDEN.search(text):
        raise ValueError(f"{path}: {label} contains invalid C body characters")
    return text.strip()


def _stub_identifier(value: str, path: Path, label: str) -> str:
    if _C_BODY_FORBIDDEN.search(value) or not _C_IDENTIFIER.fullmatch(value):
        raise ValueError(f"{path}: {label} must be a valid C identifier")
    return value


def _stub_attribute(element, name: str, path: Path, label: str) -> str:
    value = element.get(name)
    if value is None or not value.strip():
        raise ValueError(f"{path}: {label} is required")
    return _stub_identifier(value, path, label)


def _parse_stub_file(path: Path) -> tuple[_StubSpec, ...]:
    try:
        root = ET.parse(path).getroot()
    except (ET.ParseError, OSError) as error:
        raise ValueError(f"{path}: cannot parse stub XML: {error}") from error
    if root.tag != "function_analysis":
        raise ValueError(f"{path}: root tag must be function_analysis")
    called = root.find("called_functions")
    if called is None:
        raise ValueError(f"{path}: called_functions is required")
    try:
        expected_count = int(_attribute(called, "count", path))
    except ValueError as error:
        raise ValueError(f"{path}: called_functions count must be an integer") from error
    functions = list(called)
    if expected_count != len(functions):
        raise ValueError(f"{path}: declares {expected_count} called functions but contains {len(functions)}")
    stubs = []
    for function in functions:
        if function.tag != "function":
            raise ValueError(f"{path}: called_functions child must be function")
        name = _stub_attribute(function, "name", path, "stub name")
        returns = []
        arguments = []
        for child in function:
            if child.tag == "return":
                returns.append(child)
            elif child.tag == "arguments":
                arguments.append(child)
            else:
                raise ValueError(f"{path}: function child must be return or arguments")
        if len(returns) > 1 or len(arguments) > 1:
            raise ValueError(f"{path}: duplicate stub return or arguments")
        return_type = ""
        return_name = ""
        if returns:
            return_element = returns[0]
            type_element = return_element.find("type")
            name_element = return_element.find("name")
            if type_element is None or name_element is None:
                raise ValueError(f"{path}: stub return definition is incomplete")
            return_type = _stub_text(type_element, path, "stub return type")
            return_name = _stub_identifier(_stub_text(name_element, path, "stub return name"), path,
                                           "stub return name")
        else:
            warnings.warn(f"{path}: stub {name} is missing return data", RuntimeWarning)
        argument_types = []
        if arguments:
            for argument in arguments[0]:
                if argument.tag != "argument":
                    raise ValueError(f"{path}: arguments child must be argument")
                type_element = argument.find("type")
                name_element = argument.find("name")
                if type_element is None or name_element is None:
                    raise ValueError(f"{path}: stub argument definition is incomplete")
                argument_types.append(_stub_text(type_element, path, "stub argument type"))
                _stub_identifier(_stub_text(name_element, path, "stub argument name"), path, "stub argument name")
        stubs.append(_StubSpec(name, return_type, return_name, tuple(argument_types)))
    return tuple(stubs)


def _stub_catalog(paths: Sequence[Path]) -> tuple[_StubSpec, ...]:
    catalog = []
    by_name = {}
    for path in paths:
        for stub in _parse_stub_file(path):
            previous = by_name.get(stub.name)
            if previous is None:
                by_name[stub.name] = stub
                catalog.append(stub)
            elif previous != stub:
                raise ValueError(f"{path}: conflicting stub definition for {stub.name}")
    return tuple(catalog)


def _is_void(return_type: str) -> bool:
    words = return_type.split()
    while words and words[0] in _STORAGE_QUALIFIERS:
        words.pop(0)
    return words == ["void"]


def _stub_body(stub: _StubSpec, value: _ValueSpec) -> str:
    if not stub.return_type or _is_void(stub.return_type):
        return ""
    for argument in value.arguments:
        if argument.name == stub.return_name or argument.name.startswith(f"{stub.return_name}.") or argument.name.startswith(f"{stub.return_name}["):
            return f"extern {stub.return_type} {stub.return_name};\rreturn {stub.return_name};\r"
    return f"{stub.return_type} temp = {{0}};\rreturn temp;\r"


def _expanded_names(name: str) -> tuple[str, ...]:
    dimensions = []
    for match in _DIMENSION.finditer(name):
        text = match.group(1)
        if not text.isdecimal() or int(text) <= 0:
            raise ValueError(f"invalid array dimension in {name}")
        dimensions.append(int(text))
    remainder = _DIMENSION.sub("", name)
    if "[" in remainder or "]" in remainder:
        raise ValueError(f"invalid array dimension in {name}")
    if not dimensions:
        return (name,)
    ranges = []
    for size in dimensions:
        ranges.append(range(size))
    expanded = []
    for indexes in product(*ranges):
        current = iter(indexes)
        expanded.append(_DIMENSION.sub(lambda match: f"[{next(current)}]", name))
    return tuple(expanded)


def _pointer_root(name: str) -> str:
    return name.split(".", 1)[0].split("[", 1)[0]


def _pointer_groups(arguments: Sequence[_ValueArgument]) -> dict[str, list[_ValueArgument]]:
    groups = {}
    for argument in arguments:
        if argument.pointer:
            groups.setdefault(_pointer_root(argument.name), []).append(argument)
    return groups


def _validate_pointer_groups(groups: dict[str, list[_ValueArgument]], row_count: int) -> None:
    for root, arguments in groups.items():
        formats = set()
        for argument in arguments:
            formats.add(argument.fmt)
        if len(formats) != 1:
            raise ValueError(f"pointer {root} has incompatible formats")
        for row in range(row_count):
            null_states = set()
            for argument in arguments:
                if argument.direction != "out":
                    null_states.add(argument.values[row] == "NULL")
            if len(null_states) > 1:
                raise ValueError(f"pointer {root} has inconsistent NULL state")


def _storage_name(root: str, name: str) -> str:
    return f"tessy_{root}{name[len(root):]}"


def _pointer_values(argument: _ValueArgument, row_count: int, marker: str) -> tuple[str, ...]:
    if marker == "o":
        return ("",) * row_count
    values = []
    for value in argument.values:
        values.append("0" if value == "NULL" else value)
    return tuple(values)


def _plan_columns(spec: _ValueSpec) -> tuple[_Column, ...]:
    groups = _pointer_groups(spec.arguments)
    _validate_pointer_groups(groups, spec.row_count)
    columns = []
    address_roots = set()
    storage_roots = set()
    for marker, directions in (("i", {"in", "inout"}), ("o", {"out", "inout"})):
        handled_roots = set()
        for argument in spec.arguments:
            if argument.direction not in directions:
                continue
            if not argument.pointer:
                values = argument.values if marker == "i" else ("",) * spec.row_count
                for name in _expanded_names(argument.name):
                    columns.append(_Column(name, marker, values))
                continue
            root = _pointer_root(argument.name)
            if root in handled_roots:
                continue
            handled_roots.add(root)
            group = groups[root]
            if root not in address_roots:
                values = ("",) * spec.row_count
                for candidate in group:
                    if candidate.direction in ("in", "inout"):
                        address_values = []
                        for value in candidate.values:
                            address_values.append("NULL" if value == "NULL" else f"&tessy_{root}")
                        values = tuple(address_values)
                        break
                columns.append(_Column(root, marker, values))
                address_roots.add(root)
            for candidate in group:
                if candidate.direction not in directions:
                    continue
                values = _pointer_values(candidate, spec.row_count, marker)
                for name in _expanded_names(candidate.name):
                    declaration = candidate.fmt if root not in storage_roots else marker
                    columns.append(_Column(_storage_name(root, name), declaration, values))
                    storage_roots.add(root)
    return tuple(columns)


def _extension(total_variables: int) -> str:
    return ".xlsx" if total_variables <= XLSX_VARIABLE_LIMIT else ".xls"


def _partition_columns(columns: Sequence[_Column]) -> tuple[tuple[_Column, ...], ...]:
    parts = []
    for start in range(0, len(columns), XLS_VARIABLES_PER_SHEET):
        parts.append(tuple(columns[start:start + XLS_VARIABLES_PER_SHEET]))
    return tuple(parts)


def _manifest_suffix(path: Path, module: str, kind: str) -> str:
    prefix = f"{kind}_{module}_"
    if path.suffix != ".xml" or not path.name.startswith(prefix) or not path.stem[len(prefix):]:
        raise RuntimeError(f"{path.name} is not a {kind}_{module}_*.xml manifest")
    return path.stem[len(prefix):]


def _manifest_paths(paths: Sequence[Path], module: str, kind: str) -> dict[str, Path]:
    manifest = {}
    for source in paths:
        path = Path(source)
        suffix = _manifest_suffix(path, module, kind)
        if suffix in manifest:
            raise RuntimeError(f"duplicate manifest suffix {suffix}")
        manifest[suffix] = path
    return manifest


def _copy_style(source, target) -> None:
    target._style = copy(source._style)


def _extend_dimensions(sheet, first_row: int, last_row: int, first_column: int, last_column: int,
                       anchor_row: int, anchor_column: int) -> None:
    existing_rows = sheet.max_row
    existing_columns = sheet.max_column
    anchor_height = sheet.row_dimensions[anchor_row].height
    anchor_width = sheet.column_dimensions[get_column_letter(anchor_column)].width
    for row in range(max(first_row, existing_rows + 1), last_row + 1):
        sheet.row_dimensions[row].height = anchor_height
    for column in range(max(first_column, existing_columns + 1), last_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = anchor_width


def _write_cell(sheet, row: int, column: int, value, anchor_row: int, anchor_column: int) -> None:
    cell = sheet.cell(row, column)
    _copy_style(sheet.cell(anchor_row, anchor_column), cell)
    cell.value = value


def _validate_template(path: Path) -> None:
    if not path.is_file():
        raise RuntimeError(f"template does not exist: {path}")
    try:
        workbook = load_workbook(path)
    except Exception as error:
        raise RuntimeError(f"cannot load template {path}: {error}") from error
    try:
        for name in _REQUIRED_SHEETS:
            count = workbook.sheetnames.count(name)
            if count == 0:
                raise RuntimeError(f"template is missing required sheet {name}")
            if count != 1:
                raise RuntimeError(f"template has duplicate required sheet {name}")
    finally:
        workbook.close()


def _validate_limits(spec: _ValueSpec, columns: Sequence[_Column], stubs: Sequence[_StubSpec]) -> None:
    if len(columns) > XLSX_VARIABLE_LIMIT:
        raise ValueError(f"{spec.path}: has {len(columns)} variables; xlsx limit is {XLSX_VARIABLE_LIMIT}")
    testcase_limit = XLSX_ROW_LIMIT - 4
    if spec.row_count > testcase_limit:
        raise ValueError(f"{spec.path}: has {spec.row_count} test cases; xlsx limit is {testcase_limit}")
    if len(stubs) > XLSX_ROW_LIMIT - 4:
        raise ValueError(f"module has {len(stubs)} stubs; xlsx limit is {XLSX_ROW_LIMIT - 4}")


class ExcelWriter:
    def __init__(self, template_path=DEFAULT_TEMPLATE_PATH):
        self.template_path = Path(template_path)

    def write_module(self, module: str, value_paths: Sequence[Path], stub_paths: Sequence[Path],
                     output_directory: Path) -> list[Path]:
        values = _manifest_paths(value_paths, module, "value")
        stubs = _manifest_paths(stub_paths, module, "stub")
        for suffix, path in values.items():
            if suffix not in stubs:
                raise RuntimeError(f"{path.name} requires stub_{module}_{suffix}.xml")
        for suffix, path in stubs.items():
            if suffix not in values:
                raise RuntimeError(f"{path.name} requires value_{module}_{suffix}.xml")
        plans = []
        names = set()
        for suffix, path in values.items():
            spec = _parse_value(path)
            columns = _plan_columns(spec)
            destination = Path(output_directory) / f"{spec.function}.xlsx"
            if destination.name in names:
                raise RuntimeError(f"duplicate output name {destination.name}")
            names.add(destination.name)
            plans.append((spec, columns, destination))
        catalog = _stub_catalog(tuple(stubs.values()))
        for spec, columns, destination in plans:
            _validate_limits(spec, columns, catalog)
        _validate_template(self.template_path)
        output_directory = Path(output_directory)
        try:
            output_directory.mkdir(parents=True, exist_ok=True)
        except OSError as error:
            raise RuntimeError(f"cannot create output directory {output_directory}: {error}") from error
        failures = []
        written = []
        for spec, columns, destination in plans:
            try:
                self._write_workbook(spec, columns, catalog, destination)
            except Exception as error:
                failures.append(f"{spec.function} ({destination}): {error}")
            else:
                written.append(destination)
        if failures:
            raise RuntimeError("failed to write workbooks: " + "; ".join(failures))
        return written

    def _write_workbook(self, spec: _ValueSpec, columns: Sequence[_Column], stubs: Sequence[_StubSpec],
                        destination: Path) -> None:
        temporary = None
        workbook = None
        try:
            workbook = load_workbook(self.template_path)
            values = workbook["Values"]
            properties = workbook["Properties"]
            stub_sheet = workbook["Stubs"]
            _extend_dimensions(values, 1, spec.row_count + 2, 1, len(columns) + 1, 3, 2)
            for index, column in enumerate(columns, 2):
                _write_cell(values, 1, index, column.name, 1, 2)
                _write_cell(values, 2, index, column.declaration, 2, 2)
                for row, value in enumerate(column.values, 3):
                    _write_cell(values, row, index, value or None, 3, 2)
            for row in range(3, spec.row_count + 3):
                _write_cell(values, row, 1, f"tc{row - 2}.1", 3, 2)
            _extend_dimensions(properties, 5, spec.row_count + 4, 1, 5, 5, 1)
            for row in range(5, spec.row_count + 5):
                testcase = row - 4
                _write_cell(properties, row, 1, testcase, 5, 1)
                _write_cell(properties, row, 2, 0, 5, 2)
                _write_cell(properties, row, 5,
                            f"Test Case ID: UT_{spec.function}_{testcase}\rTest Case Type: Generated Test\r"
                            f"Overview: Automatically generated testcase {testcase} for {spec.function}()", 5, 5)
            _extend_dimensions(stub_sheet, 5, len(stubs) + 4, 1, 4, 5, 1)
            for row, stub in enumerate(stubs, 5):
                _write_cell(stub_sheet, row, 1, 0, 5, 1)
                _write_cell(stub_sheet, row, 2, 0, 5, 2)
                _write_cell(stub_sheet, row, 3, stub.name, 5, 3)
                _write_cell(stub_sheet, row, 4, _stub_body(stub, spec) or None, 5, 4)
            with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".tmp.xlsx", delete=False) as file:
                temporary = Path(file.name)
            workbook.save(temporary)
            workbook.close()
            workbook = None
            checked = load_workbook(temporary)
            checked.close()
            os.replace(temporary, destination)
            temporary = None
        finally:
            if workbook is not None:
                workbook.close()
            if temporary is not None:
                temporary.unlink(missing_ok=True)
