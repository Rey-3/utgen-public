"""Read workbook cells with archive-size limits and checks for unsupported formulas."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
import os
from pathlib import Path, PurePosixPath
import posixpath
import re
from types import MappingProxyType
from typing import Any, Iterable, Iterator, Mapping
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook


MAX_ARCHIVE_ENTRIES = 4096
MAX_ARCHIVE_MEMBER_BYTES = 64 * 1024 * 1024
MAX_ARCHIVE_TOTAL_BYTES = 256 * 1024 * 1024
MAX_COMPRESSION_RATIO = 1000

_CELL_REFERENCE_RE = re.compile(r"^\$?([A-Z]+)\$?([1-9][0-9]*)$")
_DIRECTION_RE = re.compile(r"^(i|o)(?:\((.+)\))?$")
_ARRAY_INDEX_RE = re.compile(r"\[([0-9]+)\]")


@dataclass(frozen=True)
class CellValue:
    """A workbook cell retaining its lexical and semantic value."""

    kind: str
    raw_text: str | None
    semantic_value: object | None

    def __post_init__(self) -> None:
        if self.kind not in {"value", "explicit_empty", "missing"}:
            raise ValueError(f"unsupported cell kind: {self.kind!r}")
        if self.raw_text is not None and not isinstance(self.raw_text, str):
            raise TypeError("raw_text must be a string or None")


@dataclass(frozen=True)
class ExcelColumnKey:
    """Immutable identity for one Values column and array element."""

    canonical_name: str
    channel: str
    element_path: tuple[int, ...]
    occurrence: int

    def __post_init__(self) -> None:
        if not isinstance(self.canonical_name, str) or not self.canonical_name:
            raise ValueError("canonical_name must be a non-empty string")
        if self.channel not in {"input", "output"}:
            raise ValueError(f"unsupported column channel: {self.channel!r}")
        if type(self.element_path) is not tuple or any(
            type(index) is not int or index < 0 for index in self.element_path
        ):
            raise TypeError("element_path must be a tuple of non-negative integers")
        if type(self.occurrence) is not int or self.occurrence < 0:
            raise TypeError("occurrence must be a non-negative integer")


@dataclass(frozen=True)
class DirectionToken:
    direction: str
    channel: str
    raw: str
    raw_type_hint: str | None


class WorkbookAdapterError(ValueError):
    """Base class for rejected or unreadable workbooks."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        sheet: str | None = None,
        coordinate: str | None = None,
    ) -> None:
        self.code = code
        self.sheet = sheet
        self.coordinate = coordinate
        location = ""
        if sheet is not None:
            location = f" [{sheet}"
            if coordinate is not None:
                location += f"!{coordinate}"
            location += "]"
        super().__init__(f"{code}{location}: {message}")


class WorkbookSecurityError(WorkbookAdapterError):
    def __init__(self, message: str, **kwargs: Any) -> None:
        super().__init__("UNSAFE_XLSX", message, **kwargs)


class WorkbookValidationError(WorkbookAdapterError):
    pass


class MissingSheetError(WorkbookValidationError):
    def __init__(self, sheet_name: str) -> None:
        super().__init__("MISSING_REQUIRED_SHEET", f"required sheet {sheet_name!r} is missing")


class UnsupportedFormulaError(WorkbookValidationError):
    def __init__(self, sheet: str, coordinate: str) -> None:
        super().__init__(
            "UNSUPPORTED_FORMULA_CELL",
            "formulas are not supported in active TESSY data regions",
            sheet=sheet,
            coordinate=coordinate,
        )


class GeneralIdentityError(WorkbookValidationError):
    def __init__(self, field: str, expected: str, actual: str) -> None:
        super().__init__(
            "GENERAL_IDENTITY_MISMATCH",
            f"General {field} does not match expected {expected!r}; got {actual!r}",
            sheet="General",
        )


@dataclass(frozen=True)
class PhysicalWorksheet:
    """Immutable index of cells physically present in a worksheet XML part."""

    title: str
    physical_cells: Mapping[tuple[int, int], Any]

    def __post_init__(self) -> None:
        if type(self.title) is not str or not self.title:
            raise ValueError("physical worksheet title must be non-empty")
        normalized: dict[tuple[int, int], Any] = {}
        for key, cell in self.physical_cells.items():
            if (
                not isinstance(key, tuple)
                or len(key) != 2
                or type(key[0]) is not int
                or type(key[1]) is not int
                or key[0] < 1
                or key[1] < 1
            ):
                raise ValueError("physical cell coordinates must be positive (row, column) tuples")
            normalized[key] = cell
        object.__setattr__(self, "physical_cells", MappingProxyType(normalized))

    @property
    def used_cells(self) -> Mapping[tuple[int, int], Any]:
        return self.physical_cells

    def iter_used_cells(self) -> Iterator[Any]:
        return iter(self.physical_cells.values())

    iter_physical_cells = iter_used_cells

    def cell_at(self, row: int, column: int) -> Any | None:
        return self.physical_cells.get((row, column))


@dataclass(frozen=True)
class PhysicalWorkbook:
    worksheets: tuple[PhysicalWorksheet, ...]

    def __post_init__(self) -> None:
        if type(self.worksheets) is not tuple:
            raise TypeError("physical workbook worksheets must be a tuple")
        if not all(isinstance(item, PhysicalWorksheet) for item in self.worksheets):
            raise TypeError("physical workbook worksheets must contain PhysicalWorksheet records")
        names = [item.title for item in self.worksheets]
        if len(names) != len(set(names)):
            raise ValueError("physical workbook sheet names must be unique")

    @property
    def sheetnames(self) -> tuple[str, ...]:
        return tuple(item.title for item in self.worksheets)

    def __getitem__(self, title: str) -> PhysicalWorksheet:
        for sheet in self.worksheets:
            if sheet.title == title:
                return sheet
        raise KeyError(title)

    def iter_used_cells(self) -> Iterator[tuple[str, Any]]:
        for sheet in self.worksheets:
            for cell in sheet.iter_used_cells():
                yield sheet.title, cell


def _column_number(value: str) -> int:
    result = 0
    for character in value:
        result = result * 26 + ord(character) - ord("A") + 1
    return result


def _cell_reference(value: object) -> tuple[int, int]:
    if not isinstance(value, str):
        raise WorkbookValidationError("INVALID_CELL_REFERENCE", "worksheet cell has no coordinate")
    match = _CELL_REFERENCE_RE.fullmatch(value)
    if match is None:
        raise WorkbookValidationError("INVALID_CELL_REFERENCE", f"invalid worksheet cell coordinate {value!r}")
    return int(match.group(2), 10), _column_number(match.group(1))


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _workbook_sheet_parts(archive: zipfile.ZipFile) -> dict[str, str]:
    try:
        workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
        rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    except (KeyError, ET.ParseError, OSError) as exc:
        raise WorkbookValidationError("INVALID_XLSX", "workbook sheet metadata is not readable") from exc
    relationships = {
        element.attrib.get("Id"): element.attrib.get("Target")
        for element in rels_root
        if _local_name(element.tag) == "Relationship"
        and element.attrib.get("Id")
        and element.attrib.get("Target")
    }
    result: dict[str, str] = {}
    for element in workbook_root.iter():
        if _local_name(element.tag) != "sheet":
            continue
        name = element.attrib.get("name")
        relationship_id = next((value for key, value in element.attrib.items() if key.rsplit("}", 1)[-1] == "id"), None)
        target = relationships.get(relationship_id)
        if not name or not target:
            raise WorkbookValidationError("INVALID_XLSX", "workbook sheet relationship is incomplete")
        part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
        if part not in archive.namelist() or not part.startswith("xl/"):
            raise WorkbookValidationError("INVALID_XLSX", f"worksheet part {part!r} is missing or unsafe")
        result[name] = part
    return result


def _physicalize_loaded_workbook(workbook: Any, raw: bytes) -> PhysicalWorkbook:
    try:
        with zipfile.ZipFile(BytesIO(raw), "r") as archive:
            parts = _workbook_sheet_parts(archive)
            physical_sheets: list[PhysicalWorksheet] = []
            for worksheet in getattr(workbook, "worksheets", ()):
                title = str(getattr(worksheet, "title", ""))
                part = parts.get(title)
                if part is None:
                    raise WorkbookValidationError("INVALID_XLSX", f"worksheet {title!r} has no workbook XML relationship")
                cells: dict[tuple[int, int], Any] = {}
                try:
                    with archive.open(part, "r") as stream:
                        for _event, element in ET.iterparse(stream, events=("end",)):
                            local_name = _local_name(element.tag)
                            if local_name == "row":
                                element.clear()
                                continue
                            if local_name != "c":
                                continue
                            reference = element.attrib.get("r")
                            key = _cell_reference(reference)
                            if key in cells:
                                raise WorkbookValidationError(
                                    "DUPLICATE_PHYSICAL_CELL",
                                    f"worksheet {title!r} contains duplicate cell {reference!r}",
                                )
                            try:
                                cells[key] = worksheet[reference]
                            except Exception as exc:
                                raise WorkbookValidationError(
                                    "INVALID_CELL_REFERENCE",
                                    f"worksheet {title!r} cell {reference!r} cannot be decoded",
                                ) from exc
                            element.clear()
                except WorkbookValidationError:
                    raise
                except (ET.ParseError, KeyError, OSError) as exc:
                    raise WorkbookValidationError("INVALID_XLSX", f"worksheet {title!r} XML is not readable") from exc
                physical_sheets.append(PhysicalWorksheet(title, cells))
            return PhysicalWorkbook(tuple(physical_sheets))
    except WorkbookValidationError:
        raise
    except (OSError, zipfile.BadZipFile) as exc:
        raise WorkbookValidationError("INVALID_XLSX", "cannot build physical workbook index") from exc


def physicalize_workbook(workbook: Any, *, raw: bytes | bytearray | None = None) -> PhysicalWorkbook:
    if isinstance(workbook, PhysicalWorkbook):
        return workbook
    if raw is not None:
        return _physicalize_loaded_workbook(workbook, bytes(raw))
    worksheets = getattr(workbook, "worksheets", None)
    if worksheets is None:
        raise WorkbookValidationError("UNSUPPORTED_WORKBOOK_STORAGE", "workbook has no worksheet collection or physical-cell protocol")
    return PhysicalWorkbook(tuple(physicalize_worksheet(item, index=index) for index, item in enumerate(worksheets)))


def physicalize_worksheet(worksheet: Any, *, index: int = 0) -> PhysicalWorksheet:
    if isinstance(worksheet, PhysicalWorksheet):
        return worksheet
    cells = getattr(worksheet, "_cells", None)
    if not isinstance(cells, Mapping):
        raise WorkbookValidationError(
            "UNSUPPORTED_WORKBOOK_STORAGE",
            "direct worksheet is neither a PhysicalWorksheet nor an openpyxl worksheet "
            "with a physical map",
        )
    return PhysicalWorksheet(str(getattr(worksheet, "title", index)), dict(cells))


def _cell_map(ws: Any) -> Mapping[tuple[int, int], Any]:
    cells = getattr(ws, "physical_cells", None)
    if isinstance(cells, Mapping):
        return cells
    iterator = getattr(ws, "iter_used_cells", None)
    if callable(iterator):
        result: dict[tuple[int, int], Any] = {}
        for cell in iterator():
            row, column = getattr(cell, "row", None), getattr(cell, "column", None)
            if type(row) is not int or type(column) is not int or row < 1 or column < 1:
                raise WorkbookValidationError("UNSUPPORTED_WORKBOOK_STORAGE", "physical-cell iterator yielded an invalid coordinate")
            key = (row, column)
            if key in result:
                raise WorkbookValidationError("DUPLICATE_PHYSICAL_CELL", f"physical-cell iterator yielded duplicate coordinate {key!r}")
            result[key] = cell
        return result
    if isinstance(getattr(ws, "_cells", None), Mapping):
        return physicalize_worksheet(ws).physical_cells
    raise WorkbookValidationError("UNSUPPORTED_WORKBOOK_STORAGE", "worksheet does not expose the project physical used-cell map protocol")


def _cell_at(ws: Any, row: int, column: int) -> Any | None:
    return _cell_map(ws).get((row, column))


def _is_formula(cell: Any | None) -> bool:
    if cell is None:
        return False
    if getattr(cell, "data_type", None) == "f":
        return True
    value = getattr(cell, "value", None)
    return isinstance(value, str) and value.startswith("=") and getattr(cell, "data_type", None) != "s"


def _cell_value(cell: Any | None) -> CellValue:
    if cell is None:
        return CellValue("missing", None, None)
    value = getattr(cell, "value", None)
    if value is None or value == "":
        return CellValue("explicit_empty", "", None)
    if isinstance(value, (datetime, date, time)):
        raw = value.isoformat()
        return CellValue("value", raw, raw)
    if isinstance(value, bool):
        return CellValue("value", "TRUE" if value else "FALSE", value)
    if isinstance(value, str):
        return CellValue("value", value, value)
    if isinstance(value, (int, float)):
        return CellValue("value", str(value), value)
    raw = str(value)
    return CellValue("value", raw, raw)


def _normalize_label(value: str) -> str:
    return " ".join(value.strip().split()).casefold()


def _text_value(cell: Any | None, *, empty: str | None = None) -> str | None:
    decoded = _cell_value(cell)
    return empty if decoded.kind == "missing" else decoded.raw_text


def _require_text(cell: Any | None, *, sheet: str, field: str, coordinate: str | None = None) -> str:
    text = _text_value(cell)
    if text is None or not text.strip():
        raise WorkbookValidationError("MISSING_REQUIRED_VALUE", f"{field} must be non-empty", sheet=sheet, coordinate=coordinate)
    return text


def _parse_integer(cell: Any | None, *, sheet: str, field: str, row: int, allow_zero: bool = True) -> int:
    text = _text_value(cell)
    if text is None or not text.strip() or re.fullmatch(r"[0-9]+", text.strip()) is None:
        raise WorkbookValidationError("INVALID_INTEGER_CELL", f"{field} must be an unsigned integer", sheet=sheet, coordinate=f"{field}{row}")
    number = int(text.strip(), 10)
    if not allow_zero and number == 0:
        raise WorkbookValidationError("INVALID_INTEGER_CELL", f"{field} must be greater than zero", sheet=sheet)
    return number


def _active_formula_check(ws: Any, *, sheet: str, min_row: int) -> None:
    for (row, _column), cell in _cell_map(ws).items():
        if row >= min_row and _is_formula(cell):
            raise UnsupportedFormulaError(sheet, cell.coordinate)


def _header_map(ws: Any, *, sheet: str, row: int, required: Iterable[str]) -> dict[str, int]:
    cells = _cell_map(ws)
    seen: dict[str, int] = {}
    for column in sorted({column for (cell_row, column) in cells if cell_row == row}):
        cell = cells[(row, column)]
        text = _text_value(cell)
        if text is None or not text.strip():
            continue
        label = _normalize_label(text)
        if label in seen:
            raise WorkbookValidationError(
                "DUPLICATE_HEADER",
                f"duplicate header {text!r}",
                sheet=sheet,
                coordinate=cell.coordinate,
            )
        seen[label] = column
    for name in required:
        if _normalize_label(name) not in seen:
            raise WorkbookValidationError(
                "MISSING_REQUIRED_HEADER",
                f"header {name!r} is missing",
                sheet=sheet,
            )
    return seen


def _row_has_payload(cells: Mapping[tuple[int, int], Any], row: int, columns: Iterable[int]) -> bool:
    for column in columns:
        cell = cells.get((row, column))
        if cell is not None and (
            getattr(cell, "value", None) is not None
            or getattr(cell, "data_type", None) == "f"
        ):
            return True
    return False


def parse_direction_token(value: Any) -> DirectionToken:
    if not isinstance(value, str):
        raise WorkbookValidationError(
            "INVALID_DIRECTION",
            f"direction token must be text, got {type(value).__name__}",
        )
    match = _DIRECTION_RE.fullmatch(value)
    if match is None:
        raise WorkbookValidationError(
            "INVALID_DIRECTION",
            f"invalid Values direction/type token {value!r}; "
            "expected i, o, i(type), or o(type)",
        )
    channel = "input" if match.group(1) == "i" else "output"
    direction = "in" if channel == "input" else "out"
    return DirectionToken(direction, channel, value, match.group(2))


parse_direction = parse_direction_token


def split_array_column_name(name: str) -> tuple[str, tuple[int, ...]]:
    if not isinstance(name, str) or not name.strip():
        raise WorkbookValidationError(
            "MISSING_COLUMN_NAME",
            "Values variable name must be non-empty text",
        )
    paths = tuple(int(item) for item in _ARRAY_INDEX_RE.findall(name))
    base = _ARRAY_INDEX_RE.sub("", name)
    if not base:
        raise WorkbookValidationError(
            "INVALID_COLUMN_NAME",
            f"Values variable name {name!r} has no base",
        )
    return base, paths


def load_xls_workbook(raw):
    """Read XLS into the existing physical-cell parser without executing formulas."""
    import struct
    import xlrd
    from openpyxl import Workbook

    if len(raw) > MAX_ARCHIVE_TOTAL_BYTES:
        raise WorkbookSecurityError("XLS input is too large")
    workbook = Workbook()
    workbook.remove(workbook.active)
    try:
        with xlrd.open_workbook(file_contents=raw, formatting_info=True, on_demand=True) as book:
            for index, sheet in enumerate(book.sheets()):
                target = workbook.create_sheet(sheet.name)
                for row in range(sheet.nrows):
                    for column in range(sheet.ncols):
                        cell = sheet.cell(row, column)
                        if cell.ctype == xlrd.XL_CELL_EMPTY:
                            continue
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_BLANK:
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, book.datemode)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code[value]
                        elif cell.ctype == xlrd.XL_CELL_NUMBER and value.is_integer():
                            value = int(value)
                        decoded = target.cell(row + 1, column + 1, value)
                        if cell.ctype == xlrd.XL_CELL_TEXT:
                            decoded.data_type = "s"
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            decoded.data_type = "e"
                # xlrd exposes cached formula results; retain formula locations for rejection.
                book._position = book._sh_abs_posn[index]
                while True:
                    opcode, _, data = book.get_record_parts()
                    if opcode == 0x000A:  # BIFF end of sheet.
                        break
                    if opcode in (0x0006, 0x0206, 0x0406):
                        row, column = struct.unpack_from("<HH", data)
                        target.cell(row + 1, column + 1, "=XLS_FORMULA")
    except WorkbookAdapterError:
        raise
    except Exception as exc:
        raise WorkbookValidationError("INVALID_XLS", "cannot decode XLS workbook") from exc
    return workbook


def read_xlsx_bytes(value: str | os.PathLike[str] | Path | bytes | bytearray) -> bytes:
    if isinstance(value, (bytes, bytearray)):
        data = bytes(value)
    else:
        try:
            data = Path(os.fspath(value)).read_bytes()
        except (TypeError, OSError) as exc:
            raise WorkbookSecurityError("cannot read XLSX input") from exc
    if not data:
        raise WorkbookSecurityError("XLSX input is empty")
    return data


def validate_xlsx_container(data: bytes) -> None:
    if not zipfile.is_zipfile(BytesIO(data)):
        raise WorkbookSecurityError("input is not a ZIP-based XLSX workbook")
    try:
        with zipfile.ZipFile(BytesIO(data), "r") as archive:
            infos = archive.infolist()
            if len(infos) > MAX_ARCHIVE_ENTRIES:
                raise WorkbookSecurityError("XLSX contains too many ZIP members")
            total = 0
            names: set[str] = set()
            for info in infos:
                name = info.filename
                if name in names:
                    raise WorkbookSecurityError(f"XLSX contains duplicate ZIP member {name!r}")
                names.add(name)
                if "\x00" in name or "\\" in name:
                    raise WorkbookSecurityError(f"XLSX ZIP member has unsafe name {name!r}")
                path = PurePosixPath(name)
                if path.is_absolute() or ".." in path.parts:
                    raise WorkbookSecurityError(f"XLSX ZIP member has unsafe path {name!r}")
                if info.file_size < 0 or info.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise WorkbookSecurityError(f"XLSX ZIP member is too large: {name!r}")
                total += info.file_size
                if total > MAX_ARCHIVE_TOTAL_BYTES:
                    raise WorkbookSecurityError("XLSX uncompressed content is too large")
                if info.compress_size == 0 and info.file_size > 0:
                    raise WorkbookSecurityError(f"XLSX ZIP member has invalid compression size: {name!r}")
                if info.compress_size and info.file_size / info.compress_size > MAX_COMPRESSION_RATIO:
                    raise WorkbookSecurityError(f"XLSX ZIP member has an unsafe compression ratio: {name!r}")
            missing = sorted({"[Content_Types].xml", "xl/workbook.xml"} - names)
            if missing:
                raise WorkbookSecurityError(f"XLSX container is missing required member(s): {', '.join(missing)}")
    except WorkbookAdapterError:
        raise
    except (OSError, zipfile.BadZipFile, RuntimeError) as exc:
        raise WorkbookSecurityError("cannot inspect XLSX ZIP container") from exc


__all__ = [
    "CellValue", "DirectionToken", "ExcelColumnKey", "GeneralIdentityError", "MissingSheetError",
    "PhysicalWorkbook", "PhysicalWorksheet", "UnsupportedFormulaError",
    "WorkbookAdapterError", "WorkbookSecurityError", "WorkbookValidationError",
    "MAX_ARCHIVE_ENTRIES", "MAX_ARCHIVE_MEMBER_BYTES", "MAX_ARCHIVE_TOTAL_BYTES",
    "MAX_COMPRESSION_RATIO", "load_workbook", "parse_direction", "parse_direction_token",
    "physicalize_workbook", "physicalize_worksheet", "read_xlsx_bytes",
    "split_array_column_name", "validate_xlsx_container",
]
