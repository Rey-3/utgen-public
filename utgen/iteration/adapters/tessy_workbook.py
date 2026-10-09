"""TESSY workbook reader and JSON-state projection boundary."""

from __future__ import annotations

from dataclasses import dataclass, field
from hashlib import sha256
from io import BytesIO
import os
from pathlib import Path
import re
import tempfile
import time
import warnings
from typing import Any, Iterable, Mapping

from openpyxl import Workbook

from tessy_excel import (
    convertDirectionToTessyFormat,
    save_test_workbook,
    value_sheet_names,
    writeGeneralSheet,
    write_testlist_cells,
)
from ..state import _stub_matches_catalog, ordered_active_rows
from ..values import flatten_values_many, pointer_member_storage, project_generated_values
from ._tessy_workbook_io import (
    CellValue,
    DirectionToken,
    ExcelColumnKey,
    GeneralIdentityError,
    MAX_ARCHIVE_ENTRIES,
    MAX_ARCHIVE_MEMBER_BYTES,
    MAX_ARCHIVE_TOTAL_BYTES,
    MAX_COMPRESSION_RATIO,
    MissingSheetError,
    PhysicalWorkbook,
    PhysicalWorksheet,
    UnsupportedFormulaError,
    WorkbookAdapterError,
    WorkbookSecurityError,
    WorkbookValidationError,
    _cell_at,
    _cell_map,
    _cell_value,
    _active_formula_check,
    _header_map,
    _normalize_label,
    parse_direction,
    parse_direction_token,
    _parse_integer,
    _require_text,
    _row_has_payload,
    split_array_column_name,
    _text_value,
    load_workbook,
    load_xls_workbook,
    physicalize_workbook,
    physicalize_worksheet,
    read_xlsx_bytes,
    validate_xlsx_container,
)


_CASE_LABEL_RE = re.compile(r"^tc([1-9][0-9]*)\.([0-9]+)$")
_TESSY_51_RE = re.compile(r"^5\.1(?:\.[0-9]+)?$")
_FORMAL_ID_RE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_.-]*_[0-9]{3,}$")


def default_stub_body(catalog_entry, used) -> str:
    """Return the canonical body for one catalog entry in TESSY output."""
    return_type = catalog_entry["return_type"]
    if return_type == "void":
        return ""
    storage = catalog_entry["return_storage"]
    if used:
        return f"return {storage};"
    return f"{return_type} temp = {{0}};\nreturn temp;"


@dataclass(frozen=True)
class WorkbookMeta:
    project: str
    module: str
    testobject: str
    tessy_version: str
    export_date: str | None = None
    pdb_file: str | None = None
    dialect: str = "tessy_5_1"

    @property
    def project_name(self) -> str:
        return self.project

    @property
    def module_name(self) -> str:
        return self.module

    @property
    def test_object(self) -> str:
        return self.testobject

    @property
    def test_object_name(self) -> str:
        return self.testobject

    @property
    def tessy(self) -> str:
        return self.tessy_version


@dataclass(frozen=True)
class ExportColumn:
    key: ExcelColumnKey
    name: str
    direction: str
    channel: str
    raw_direction: str
    raw_type_hint: str | None
    element_path: tuple[int, ...]
    base_name: str
    source_column: int

    @property
    def canonical_name(self) -> str:
        return self.base_name

    @property
    def type_hint(self) -> str | None:
        return self.raw_type_hint

    @property
    def column(self) -> int:
        return self.source_column

    @property
    def direction_token(self) -> str:
        return self.raw_direction


@dataclass
class ExportCase:
    export_case_number: int
    export_step: int
    row_label: str
    case_name: str | None
    external_test_case_id: str | None
    description_raw: str
    input_cells: dict[ExcelColumnKey, CellValue] = field(default_factory=dict)
    output_cells: dict[ExcelColumnKey, CellValue] = field(default_factory=dict)
    properties_test_step: int | None = None
    properties_row: int | None = None
    values_row: int | None = None
    description_text: str | None = None

    @property
    def test_case(self) -> int:
        return self.export_case_number

    @property
    def test_step(self) -> int:
        return self.export_step

    @property
    def name(self) -> str | None:
        return self.case_name

    @property
    def description(self) -> str:
        return self.description_raw

    @property
    def row_id(self) -> str | None:
        return self.external_test_case_id

    @staticmethod
    def _values(cells: Mapping[ExcelColumnKey, CellValue]) -> dict[str, str]:
        return {
            key.canonical_name: (cell.raw_text if cell.raw_text is not None else "")
            for key, cell in cells.items()
        }

    @property
    def input(self) -> dict[str, str]:
        return self._values(self.input_cells)

    @property
    def output(self) -> dict[str, str]:
        return self._values(self.output_cells)

    @property
    def state_description(self) -> str:
        return self.description_text if self.description_text is not None else self.description_raw


@dataclass(frozen=True)
class ExportStub:
    name: str
    test_case: int
    test_step: int
    body_raw: str
    source_row: int | None = None

    @property
    def stub_name(self) -> str:
        return self.name

    @property
    def body(self) -> str:
        return self.body_raw


@dataclass
class WorkbookSnapshot:
    meta: WorkbookMeta
    cases: list[ExportCase]
    columns: list[ExportColumn]
    stubs: list[ExportStub]
    completeness: str
    workbook_hash: str

    @property
    def hash(self) -> str:
        return self.workbook_hash

    def case(self, row_id: str) -> ExportCase:
        for item in self.cases:
            if item.external_test_case_id == row_id:
                return item
        raise KeyError(row_id)


@dataclass(frozen=True)
class _PropertyRecord:
    case_number: int
    test_step: int
    case_name: str | None
    description_raw: str
    external_test_case_id: str | None
    description_text: str | None
    source_row: int


def _error(code: str, message: str, *, sheet: str | None = None, coordinate: str | None = None) -> None:
    raise WorkbookValidationError(code, message, sheet=sheet, coordinate=coordinate)


def _module_basename(value: str) -> str:
    return value.replace("\\", "/").rsplit("/", 1)[-1]


def _is_formal_description(description: str) -> bool:
    return any(
        ":" in line
        and _normalize_label(line.split(":", 1)[0]) in {"test case id", "description"}
        for line in description.splitlines()
    )


def _description_parts(description: str) -> tuple[str | None, str | None]:
    lines = description.split("\n")
    if _is_formal_description(description):
        if len(lines) < 2:
            _error(
                "MALFORMED_PROPERTIES_DESCRIPTION",
                "formal Description must contain Test Case ID then Description lines",
                sheet="Properties",
            )
        first = lines[0].split(":", 1) if ":" in lines[0] else ()
        second = lines[1].split(":", 1) if ":" in lines[1] else ()
        if (
            len(first) != 2
            or _normalize_label(first[0]) != "test case id"
            or len(second) != 2
            or _normalize_label(second[0]) != "description"
        ):
            _error(
                "MALFORMED_PROPERTIES_DESCRIPTION",
                "formal Description must start with Test Case ID then Description",
                sheet="Properties",
            )
        external_id = first[1].strip()
        if not external_id or not _FORMAL_ID_RE.fullmatch(external_id):
            _error(
                "MALFORMED_ROW_ID",
                f"invalid Properties Test Case ID {external_id!r}",
                sheet="Properties",
            )
        managed = {"test case id", "description", "test case type", "overview"}
        for line in lines[2:]:
            if ":" in line and _normalize_label(line.split(":", 1)[0]) in managed:
                _error(
                    "MALFORMED_PROPERTIES_DESCRIPTION",
                    "formal Description contains an extra managed header",
                    sheet="Properties",
                )
        first_description = second[1][1:] if second[1].startswith(" ") else second[1]
        description_text = "\n".join([first_description, *lines[2:]])
        return external_id, description_text

    for line in lines:
        if _normalize_label(line).startswith("test case id") and ":" not in line:
            _error("MALFORMED_ROW_ID", "Properties Test Case ID line must use 'Test Case ID: <id>'", sheet="Properties")
    return None, None


class TessyWorkbookAdapter:
    dialect = "tessy_5_1"
    mandatory_sheets = ("General", "Properties", "Values", "Stubs")

    def __init__(self, *, expected_module=None, expected_testobject=None, expected_test_object=None, expected_project=None) -> None:
        if expected_testobject is not None and expected_test_object is not None and expected_testobject != expected_test_object:
            raise ValueError("expected_testobject and expected_test_object disagree")
        self.expected_module = expected_module
        self.expected_testobject = expected_testobject if expected_testobject is not None else expected_test_object
        self.expected_project = expected_project

    def read_complete_export(
        self,
        path,
        *,
        expected_module=None,
        expected_testobject=None,
        expected_test_object=None,
        expected_project=None,
    ) -> WorkbookSnapshot:
        if expected_testobject is not None and expected_test_object is not None and expected_testobject != expected_test_object:
            raise ValueError("expected_testobject and expected_test_object disagree")
        expected_testobject = expected_testobject if expected_testobject is not None else expected_test_object
        expected_module = self.expected_module if expected_module is None else expected_module
        expected_testobject = self.expected_testobject if expected_testobject is None else expected_testobject
        expected_project = self.expected_project if expected_project is None else expected_project
        raw = read_xlsx_bytes(path)
        is_xls = raw.startswith(bytes.fromhex("D0CF11E0A1B11AE1"))
        if is_xls:
            workbook = load_xls_workbook(raw)
        else:
            validate_xlsx_container(raw)
            try:
                workbook = load_workbook(BytesIO(raw), read_only=False, data_only=False, keep_links=False)
            except Exception as exc:
                raise WorkbookValidationError("INVALID_XLSX", "cannot decode XLSX workbook") from exc
        try:
            physical = physicalize_workbook(workbook, raw=None if is_xls else raw)
            self._require_sheets(physical)
            meta = self._parse_general(physical["General"], expected_module, expected_testobject, expected_project)
            properties = self._parse_properties(physical["Properties"])
            columns, value_rows = self._parse_values_sheets(physical)
            stubs = self._parse_stubs(physical["Stubs"])
            self._validate_cross_sheet(properties, value_rows, stubs)
            cases = self._join_cases(properties, value_rows)
        finally:
            workbook.close()
        return WorkbookSnapshot(meta, cases, columns, stubs, "complete", sha256(raw).hexdigest())

    def read(self, path, **kwargs):
        return self.read_complete_export(path, **kwargs)

    def parse(self, path, **kwargs):
        return self.read_complete_export(path, **kwargs)

    def _require_sheets(self, workbook: PhysicalWorkbook) -> None:
        for name in self.mandatory_sheets:
            if name not in workbook.sheetnames:
                raise MissingSheetError(name)

    def _parse_general(self, ws, expected_module, expected_testobject, expected_project) -> WorkbookMeta:
        labels = {}
        for (row, column), cell in _cell_map(ws).items():
            text = _text_value(cell)
            if column == 1 and text and text.strip():
                label = _normalize_label(text)
                if label in labels:
                    _error("DUPLICATE_GENERAL_LABEL", f"duplicate General label {text!r}", sheet="General")
                labels[label] = row
        required = {"project": "project", "module": "module", "testobject": "testobject", "tessy version": "tessy_version"}
        values = {}
        for label, field_name in required.items():
            if label not in labels:
                _error("MISSING_GENERAL_FIELD", f"General field {label!r} is missing", sheet="General")
            row = labels[label]
            values[field_name] = _require_text(_cell_at(ws, row, 2), sheet="General", field=label, coordinate=f"B{row}")
        module = _module_basename(values["module"])
        if not module:
            _error("MISSING_GENERAL_FIELD", "General Module must identify a source file", sheet="General")
        if not _TESSY_51_RE.fullmatch(values["tessy_version"]):
            _error("UNSUPPORTED_TESSY_DIALECT", f"unsupported TESSY version {values['tessy_version']!r}", sheet="General")
        export_date = _text_value(_cell_at(ws, labels["export date"], 2)) if "export date" in labels else None
        if export_date == "":
            export_date = None
        pdb_file = _text_value(_cell_at(ws, labels["pdb file"], 2)) if "pdb file" in labels else None
        meta = WorkbookMeta(values["project"], module, values["testobject"], values["tessy_version"], export_date, pdb_file, self.dialect)
        if expected_module is not None and meta.module != _module_basename(os.fspath(expected_module)):
            raise GeneralIdentityError("Module", _module_basename(os.fspath(expected_module)), meta.module)
        if expected_testobject is not None and meta.testobject != expected_testobject:
            raise GeneralIdentityError("Testobject", expected_testobject, meta.testobject)
        if expected_project is not None and meta.project != expected_project:
            raise GeneralIdentityError("Project", expected_project, meta.project)
        return meta

    def _parse_properties(self, ws) -> dict[int, _PropertyRecord]:
        _active_formula_check(ws, sheet="Properties", min_row=4)
        headers = _header_map(ws, sheet="Properties", row=4, required=("Test Case", "Test Step", "Name", "Description"))
        cells = _cell_map(ws)
        records = {}
        ids: set[str] = set()
        candidate_rows = sorted({row for (row, _column), cell in cells.items() if row >= 5 and getattr(cell, "value", None) is not None})
        for row in candidate_rows:
            if not _row_has_payload(cells, row, headers.values()):
                continue
            case_number = _parse_integer(
                cells.get((row, headers[_normalize_label("Test Case")])),
                sheet="Properties", field="Test Case", row=row, allow_zero=False,
            )
            test_step = _parse_integer(
                cells.get((row, headers[_normalize_label("Test Step")])),
                sheet="Properties", field="Test Step", row=row,
            )
            if case_number in records:
                _error("DUPLICATE_PROPERTY_CASE", f"duplicate Properties Test Case {case_number}", sheet="Properties")
            name = _text_value(cells.get((row, headers[_normalize_label("Name")])))
            description = _text_value(cells.get((row, headers[_normalize_label("Description")]))) or ""
            external_id, description_text = _description_parts(description)
            if external_id is not None and external_id in ids:
                _error("DUPLICATE_ROW_ID", f"duplicate Properties row ID {external_id!r}", sheet="Properties")
            if external_id is not None:
                ids.add(external_id)
            records[case_number] = _PropertyRecord(
                case_number, test_step,
                name if name and name.strip() else None, description,
                external_id, description_text, row,
            )
        return records

    def _parse_values_sheets(self, workbook):
        names = value_sheet_names(workbook.sheetnames)
        expected = ["Values", *(f"Values {i}" for i in range(1, len(names)))]
        if names != expected:
            _error("MISSING_VALUES_SHEET", "Values pages must be numbered consecutively", sheet="Values")
        columns, rows, seen = [], {}, set()
        for name in names:
            row_labels = {row[5]: row[2] for row in rows.values()} if name != "Values" else None
            page_columns, page_rows = self._parse_values(workbook[name], row_labels)
            for column in page_columns:
                key = (column.base_name, column.channel, column.element_path)
                if key in seen:
                    _error("DUPLICATE_ARGUMENT_COLUMN", f"duplicate Values column {column.name!r}", sheet=name)
                seen.add(key)
            columns.extend(page_columns)
            page_rows = {row[2]: row for row in page_rows}
            if name == "Values":
                rows = page_rows
            else:
                if page_rows.keys() != rows.keys():
                    _error("MISSING_VALUES_RECORD", "all Values pages must contain the same testcase/step labels", sheet=name)
                for label, row in page_rows.items():
                    rows[label][3].update(row[3])
                    rows[label][4].update(row[4])
        self._validate_array_structure(columns)
        return columns, list(rows.values())

    def _parse_values(self, ws, row_labels=None):
        _active_formula_check(ws, sheet=ws.title, min_row=1)
        cells = _cell_map(ws)
        unlabelled = row_labels is not None and any(_text_value(cells.get((row, 1))) for row in (1, 2))
        if unlabelled:
            try:
                parse_direction_token(_text_value(cells.get((2, 1))))
            except WorkbookAdapterError:
                headings = tuple(_normalize_label(_text_value(cells.get((row, 1))) or "") for row in (1, 2))
                unlabelled = headings != ("test case", "test step") and not any(
                    _CASE_LABEL_RE.fullmatch(_text_value(cell) or "")
                    for (row, column), cell in cells.items() if row >= 3 and column == 1)
        first_column = 1 if unlabelled else 2
        columns = []
        seen = set()
        active_numbers = sorted({column for (row, column), cell in cells.items() if column >= first_column and getattr(cell, "value", None) is not None})
        for column in active_numbers:
            name_cell, direction_cell = cells.get((1, column)), cells.get((2, column))
            name, raw_direction = _text_value(name_cell), _text_value(direction_cell)
            if name is None or not name.strip():
                _error("MISSING_COLUMN_NAME", f"active Values column {column} has no variable name", sheet=ws.title)
            if raw_direction is None or not raw_direction.strip():
                _error("INVALID_DIRECTION", f"Values column {name!r} has no direction token", sheet=ws.title)
            try:
                token = parse_direction_token(raw_direction)
            except WorkbookAdapterError as exc:
                raise WorkbookValidationError(exc.code, str(exc), sheet=ws.title, coordinate=direction_cell.coordinate if direction_cell else None) from exc
            base_name, element_path = split_array_column_name(name)
            key_tuple = (base_name, token.channel, element_path)
            if key_tuple in seen:
                _error("DUPLICATE_ARGUMENT_COLUMN", f"same-name same-channel Values column {name!r} occurs more than once", sheet=ws.title)
            seen.add(key_tuple)
            key = ExcelColumnKey(name, token.channel, element_path, 0)
            columns.append(ExportColumn(key, name, token.direction, token.channel, token.raw, token.raw_type_hint, element_path, base_name, column))
        rows = []
        labels = set()
        candidate_rows = sorted({row for (row, _column), cell in cells.items() if row >= 3 and (unlabelled or getattr(cell, "value", None) is not None)})
        for row in candidate_rows:
            label_cell = None if unlabelled else cells.get((row, 1))
            label = row_labels.get(row) if unlabelled else _text_value(label_cell)
            if not (label is not None or _row_has_payload(cells, row, active_numbers)):
                continue
            if label is None or not label.strip() or _CASE_LABEL_RE.fullmatch(label) is None:
                _error(
                    "INVALID_VALUES_ROW_LABEL",
                    f"invalid Values row label {label!r}",
                    sheet=ws.title,
                    coordinate=label_cell.coordinate if label_cell else None,
                )
            if label in labels:
                _error("DUPLICATE_VALUES_ROW", f"duplicate Values row label {label!r}", sheet=ws.title)
            labels.add(label)
            match = _CASE_LABEL_RE.fullmatch(label)
            case_number, step = int(match.group(1)), int(match.group(2))
            inputs, outputs = {}, {}
            for column in columns:
                target = inputs if column.channel == "input" else outputs
                target[column.key] = _cell_value(cells.get((row, column.source_column)))
            rows.append((case_number, step, label, inputs, outputs, row))
        return columns, rows

    def _validate_array_structure(self, columns: list[ExportColumn]) -> None:
        groups = {}
        for column in columns:
            groups.setdefault((column.base_name, column.channel), []).append(column.element_path)
        for (base, channel), paths in groups.items():
            if not any(paths):
                continue
            if any(not path for path in paths) or len({len(path) for path in paths}) != 1:
                _error("INCOMPLETE_ARRAY", f"array column group {base!r}/{channel} is structurally incomplete", sheet="Values")
            ranges = [
                set(range(max(path[i] for path in paths) + 1))
                for i in range(len(paths[0]))
            ]
            expected = {tuple(indexes) for indexes in _product_ranges(ranges)}
            if set(paths) != expected:
                _error("INCOMPLETE_ARRAY", f"array column group {base!r}/{channel} is structurally incomplete", sheet="Values")

    def _parse_stubs(self, ws) -> list[ExportStub]:
        _active_formula_check(ws, sheet="Stubs", min_row=4)
        headers = _header_map(ws, sheet="Stubs", row=4, required=("Test Case", "Test Step", "Stub", "Stub Body"))
        cells = _cell_map(ws)
        rows = sorted({row for (row, _column), cell in cells.items() if row >= 5 and getattr(cell, "value", None) is not None})
        result, names = [], set()
        for row in rows:
            if not _row_has_payload(cells, row, headers.values()):
                continue
            name = _text_value(cells.get((row, headers[_normalize_label("Stub")])))
            if name is None or not name.strip():
                _error("MISSING_STUB_NAME", "active Stubs row has no Stub name", sheet="Stubs")
            if name in names:
                _error("DUPLICATE_STUB_NAME", f"duplicate module-level Stub name {name!r}", sheet="Stubs")
            names.add(name)
            test_case = _parse_integer(cells.get((row, headers[_normalize_label("Test Case")])), sheet="Stubs", field="Test Case", row=row)
            test_step = _parse_integer(cells.get((row, headers[_normalize_label("Test Step")])), sheet="Stubs", field="Test Step", row=row)
            body = _text_value(cells.get((row, headers[_normalize_label("Stub Body")]))) or ""
            result.append(ExportStub(name, test_case, test_step, body, row))
        return result

    def _validate_cross_sheet(self, properties, value_rows, stubs) -> None:
        numbers = [row[0] for row in value_rows]
        if sorted(set(numbers) - set(properties)):
            _error("MISSING_PROPERTY_RECORD", f"Values testcase(s) lack Properties records: {sorted(set(numbers) - set(properties))}")
        if sorted(set(properties) - set(numbers)):
            _error("MISSING_VALUES_RECORD", f"Properties testcase(s) lack Values rows: {sorted(set(properties) - set(numbers))}")
        if len({stub.name for stub in stubs}) != len(stubs):
            _error("DUPLICATE_STUB_NAME", "Stub names must be unique within the module")

    def _join_cases(self, properties, value_rows) -> list[ExportCase]:
        cases = []
        for case_number, step, label, inputs, outputs, values_row in value_rows:
            property_record = properties[case_number]
            cases.append(
                ExportCase(
                    case_number,
                    step,
                    label,
                    property_record.case_name,
                    property_record.external_test_case_id,
                    property_record.description_raw,
                    inputs,
                    outputs,
                    property_record.test_step,
                    property_record.source_row,
                    values_row,
                    property_record.description_text,
                )
            )
        return cases

    def write_function(self, path, module_state: dict, function_name: str, template_path=None) -> None:
        """Atomically project one JSON function into a complete TESSY workbook."""
        self.write_functions(((function_name, path),), module_state, template_path)

    def write_functions(self, paths, module_state: dict, template_path=None) -> None:
        """Validate a module once and atomically publish its requested workbooks."""
        from ..validation import validate_module

        validate_module(module_state)
        paths = tuple(paths)
        for index, (function_name, path) in enumerate(paths, 1):
            started = time.monotonic()
            self._write_function(path, module_state, function_name, template_path)
            print(f"[Excel {index}/{len(paths)}] {function_name:<60} {time.monotonic() - started:.3f}s", flush=True)

    def _write_function(self, path, module_state: dict, function_name: str, template_path=None) -> None:
        if function_name not in module_state["functions"]:
            raise WorkbookValidationError("UNKNOWN_FUNCTION", f"function {function_name!r} is not present in module state")
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        function = module_state["functions"][function_name]
        rows = ordered_active_rows(function)
        columns, projected = _project_rows(function, rows)
        stubs = _project_stubs(module_state, function)
        existing = path if path.is_file() else path.with_suffix(".xls" if path.suffix == ".xlsx" else ".xlsx")
        if existing.is_file():
            raw = read_xlsx_bytes(existing)
            if raw.startswith(bytes.fromhex("D0CF11E0A1B11AE1")):
                workbook = load_xls_workbook(raw)
            else:
                validate_xlsx_container(raw)
                workbook = load_workbook(BytesIO(raw), keep_links=False)
        else:
            workbook = _load_template(template_path)
        _ensure_sheets(workbook)
        writeGeneralSheet(workbook["General"], module_state["project"], module_state["source_file"], function_name)
        write_testlist_cells(workbook, projected, stubs, columns)
        path = path.with_suffix(".xls" if len(columns) + 1 > 16384 else ".xlsx")
        fd, temporary_name = tempfile.mkstemp(prefix=f".{path.name}.", suffix=f".tmp{path.suffix}", dir=path.parent)
        os.close(fd)
        temporary = Path(temporary_name)
        try:
            save_test_workbook(workbook, temporary)
            workbook.close()
            snapshot = self.read_complete_export(
                temporary,
                expected_module=_module_basename(module_state["source_file"]),
                expected_testobject=function_name,
                expected_project=f"tessy_{module_state['project']}",
            )
            _validate_projection(snapshot, projected, stubs, columns)
            os.replace(temporary, path)
            alternate = path.with_suffix(".xlsx" if path.suffix == ".xls" else ".xls")
            if alternate.is_file():
                alternate.rename(alternate.with_name(f"{alternate.name}.{time.time_ns()}.orphaned"))
        finally:
            try:
                workbook.close()
            except Exception:
                pass
            try:
                temporary.unlink()
            except FileNotFoundError:
                pass


def _product_ranges(ranges: list[set[int]]) -> Iterable[tuple[int, ...]]:
    if not ranges:
        yield ()
    else:
        for value in sorted(ranges[0]):
            for suffix in _product_ranges(ranges[1:]):
                yield (value, *suffix)


def _pointer_roots(function: Mapping[str, Any]) -> set[str]:
    variables = function.get("variables", {})
    return {
        name[1:]
        for name in variables
        if name.startswith("@") and name[1:] in variables
    }


def _pointer_storage_root(function: Mapping[str, Any], name: str) -> str | None:
    for root in _pointer_roots(function):
        storage = f"tessy_{root}"
        if name == storage or name.startswith(storage + ".") or name.startswith(storage + "["):
            return root
    return None


def _is_null_pointer_address(value: str) -> bool:
    token = value.strip()
    return token == "NULL" or re.fullmatch(r"0[uUlL]*", token) is not None


def _flatten_channel(function: Mapping[str, Any], flat: Mapping[str, str], channel: str, members) -> dict[str, str]:
    pointer_roots = _pointer_roots(function)
    projected: dict[str, str] = {}
    for name, value in flat.items():
        if name in members:
            storage, _ = members[name]
            pairs = ((name, "NULL" if value.strip() == "NULL" else storage),
                     (storage, "0" if value.strip() == "NULL" else value))
        elif name.startswith("@") and name[1:] in pointer_roots:
            if channel != "input":
                continue
            root = name[1:]
            target = root
            value = "NULL" if _is_null_pointer_address(value) else f"tessy_{root}"
            pairs = ((target, value),)
        else:
            root = re.split(r"[.\[]", name, maxsplit=1)[0]
            target = f"tessy_{name}" if root in pointer_roots else name
            pairs = ((target, value),)
        for target, value in pairs:
            if target in projected:
                raise WorkbookValidationError(
                    "POINTER_ALIAS_COLLISION",
                    f"TESSY projection maps more than one JSON value to {target!r}",
                )
            projected[target] = value
    return projected


def _project_rows(function: Mapping[str, Any], rows: list[tuple[str, dict]]):
    members = pointer_member_storage(function["variables"])
    storage_types = dict(members.values())
    source_rows = [row for _, row in rows] or [
        {channel: project_generated_values({}, function["variables"], channel) for channel in ("input", "output")}
    ]
    flattened = {
        channel: [_flatten_channel(function, flat, channel, members)
                  for flat in flatten_values_many((row[channel] for row in source_rows), function["variables"], channel)]
        for channel in ("input", "output")
    }
    input_flat, output_flat = flattened["input"][0], flattened["output"][0]
    columns = []
    typed_roots: set[str] = set()
    pointer_roots = _pointer_roots(function)
    parameters = set(function.get("parameters", ()))
    input_names = list(input_flat)

    def is_parameter_column(name: str) -> bool:
        if name in pointer_roots:
            return name in parameters or not parameters
        root = re.split(r"[.\[]", name, maxsplit=1)[0]
        return root in parameters and root not in pointer_roots

    ordered_input_names = [
        *(name for name in input_names if is_parameter_column(name)),
        *(name for name in input_names if not is_parameter_column(name)),
    ]
    for name in ordered_input_names:
        columns.append((name, "input", _column_direction_token(function, name, "input", typed_roots, members, storage_types)))
    for name in output_flat:
        columns.append((name, "output", _column_direction_token(function, name, "output", typed_roots, members, storage_types)))
    projected = []
    for index, (row_id, row) in enumerate(rows):
        projected.append(
            {
                "id": row_id,
                "description": row["description"],
                "input": flattened["input"][index],
                "output": flattened["output"][index],
            }
        )
    return columns, projected


def _column_direction_token(
    function: Mapping[str, Any],
    name: str,
    channel: str,
    typed_roots: set[str],
    members,
    storage_types,
) -> str:
    direction = "in" if channel == "input" else "out"
    # A type token declares synthetic storage in TESSY. Existing C interfaces
    # already have their type and array shape; redeclaring them causes conflicts.
    if not name.startswith("tessy_"):
        return convertDirectionToTessyFormat(direction, name, show_type=False)
    if name in members:
        return convertDirectionToTessyFormat(direction, name, show_type=False)
    if name in storage_types:
        root = name.split("[")[0]
        show_type = root not in typed_roots
        typed_roots.add(root)
        return convertDirectionToTessyFormat(direction, name, storage_types[name], show_type)

    pointer_roots = _pointer_roots(function)
    if channel == "input" and name in pointer_roots:
        return convertDirectionToTessyFormat(direction, name, show_type=False)

    storage_root = _pointer_storage_root(function, name)
    if storage_root is not None:
        spec = function.get("variables", {}).get(storage_root, {})
        typed_root = f"tessy_{storage_root}"
        type_hint = spec.get("type") or spec.get("format") or spec.get("fmt") or ""
        if not type_hint or typed_root in typed_roots:
            return convertDirectionToTessyFormat(direction, name, show_type=False)
        typed_roots.add(typed_root)
        return convertDirectionToTessyFormat(direction, name, str(type_hint), True)

    root = re.split(r"[.\[]", name, maxsplit=1)[0]
    spec = function.get("variables", {}).get(root, {})
    if spec.get("pointer") or spec.get("is_pointer"):
        return convertDirectionToTessyFormat(direction, name, show_type=False)
    type_hint = spec.get("type") or spec.get("format") or spec.get("fmt") or ""
    # Declare integer aliases (including uint32/uint64) using their resolved C type.
    native_type = spec.get("format") or ""
    if native_type.startswith(("unsigned ", "signed ")) or native_type in {"short", "int", "long", "long long"}:
        type_hint = native_type
    if not type_hint or root in typed_roots:
        return convertDirectionToTessyFormat(direction, name, show_type=False)
    typed_roots.add(root)
    return convertDirectionToTessyFormat(direction, name, str(type_hint), True)


def _project_stubs(state: Mapping[str, Any], function: Mapping[str, Any]) -> list[dict[str, str]]:
    result = []
    for key, entry in state.get("stub_catalog", {}).items():
        if entry["status"] != "active":
            continue
        local = function.get("stubs", {}).get(key)
        if local is not None and local["status"] == "orphaned":
            continue
        if local is not None and local["status"] == "active":
            if not _stub_matches_catalog(function, key, entry):
                warnings.warn(f"user stub {key!r} has a changed C signature; kept unchanged in JSON, excluded from Excel", stacklevel=2)
                continue
            body = local["body"]
        else:
            body = default_stub_body(
                entry,
                key in function.get("used_stub_keys", []),
            )
        result.append({"name": key, "body": body})
    return result


def _load_template(template_path):
    if template_path is None:
        candidate = Path(__file__).resolve().parents[3] / "template.xlsx"
        if candidate.exists():
            template_path = candidate
    elif not Path(template_path).exists():
        raise WorkbookValidationError(
            "MISSING_TEMPLATE",
            f"explicit TESSY template does not exist: {template_path}",
        )
    if template_path is not None:
        return load_workbook(template_path)
    workbook = Workbook()
    workbook.active.title = "General"
    for name in ("Properties", "User Code", "Stubs", "Call Trace", "Values", "Fault Injections"):
        workbook.create_sheet(name)
    for row, label, value in (
        (9, "Project", ""),
        (10, "Module", ""),
        (11, "Testobject", ""),
        (13, "Export Date", ""),
        (15, "Tessy Version", "5.1.0"),
    ):
        workbook["General"].cell(row, 1).value = label
        workbook["General"].cell(row, 2).value = value
    for sheet, headers, row in (
        ("Properties", ("Test Case", "Test Step", "Repeat Count", "Name", "Description"), 4),
        ("Stubs", ("Test Case", "Test Step", "Stub", "Stub Body"), 4),
    ):
        for index, value in enumerate(headers, 1):
            workbook[sheet].cell(row, index).value = value
    return workbook


def _ensure_sheets(workbook) -> None:
    for name in TessyWorkbookAdapter.mandatory_sheets:
        if name not in workbook.sheetnames:
            workbook.create_sheet(name)
    for sheet, headers, row in (
        ("Properties", ("Test Case", "Test Step", "Repeat Count", "Name", "Description"), 4),
        ("Stubs", ("Test Case", "Test Step", "Stub", "Stub Body"), 4),
    ):
        target = workbook[sheet]
        for index, value in enumerate(headers, 1):
            if target.cell(row, index).value is None:
                target.cell(row, index).value = value


def _validate_projection(snapshot: WorkbookSnapshot, rows, stubs, columns) -> None:
    expected_ids = [row["id"] for row in rows]
    actual_ids = [case.external_test_case_id for case in snapshot.cases]
    if actual_ids != expected_ids:
        raise WorkbookValidationError("PROJECTION_MISMATCH", f"Properties/Values row IDs differ: expected {expected_ids!r}, got {actual_ids!r}")
    expected_columns = [(column[0], column[1]) for column in columns]
    actual_columns = [(column.name, column.channel) for column in snapshot.columns]
    if actual_columns != expected_columns:
        raise WorkbookValidationError("PROJECTION_MISMATCH", "Values columns differ from JSON state")
    for case, expected in zip(snapshot.cases, rows):
        if case.state_description != expected["description"] or case.input != expected["input"] or case.output != expected["output"]:
            raise WorkbookValidationError("PROJECTION_MISMATCH", f"Values/Properties differ for row {expected['id']!r}")
    if [(stub.name, stub.body_raw) for stub in snapshot.stubs] != [(stub["name"], stub["body"]) for stub in stubs]:
        raise WorkbookValidationError("PROJECTION_MISMATCH", "Stubs differ from JSON state")


def workbook_to_function_snapshot(path, **kwargs) -> WorkbookSnapshot:
    return TessyWorkbookAdapter().read_complete_export(path, **kwargs)


def read_complete_export(path, **kwargs) -> WorkbookSnapshot:
    return TessyWorkbookAdapter().read_complete_export(path, **kwargs)


def read_workbook(path, **kwargs) -> WorkbookSnapshot:
    return read_complete_export(path, **kwargs)


WorkbookAdapter = TessyWorkbookAdapter
read = read_complete_export
parse_workbook = read_complete_export


__all__ = [
    "CellValue", "DirectionToken", "ExcelColumnKey", "ExportCase", "ExportColumn",
    "ExportStub", "GeneralIdentityError", "MissingSheetError", "PhysicalWorkbook",
    "PhysicalWorksheet", "TessyWorkbookAdapter", "UnsupportedFormulaError",
    "WorkbookAdapter", "WorkbookAdapterError", "WorkbookMeta", "WorkbookSecurityError",
    "WorkbookSnapshot", "WorkbookValidationError", "MAX_ARCHIVE_ENTRIES",
    "MAX_ARCHIVE_MEMBER_BYTES", "MAX_ARCHIVE_TOTAL_BYTES", "MAX_COMPRESSION_RATIO",
    "parse_direction", "parse_direction_token", "parse_workbook", "read",
    "read_complete_export", "read_workbook", "workbook_to_function_snapshot",
    "physicalize_workbook", "physicalize_worksheet", "split_array_column_name",
    "default_stub_body",
]
